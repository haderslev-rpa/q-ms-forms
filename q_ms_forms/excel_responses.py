"""Læsning og sletning af svarrækker i en delt Forms-Excel-fil."""

from __future__ import annotations

import base64
import io
import logging
import time
from dataclasses import dataclass
from typing import Any

import requests
from automation_server_client import AutomationServer, Credential
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

logger = logging.getLogger(__name__)
GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
GRAPH_SCOPE = "https://graph.microsoft.com/.default"
DEFAULT_TIMEOUT_SECONDS = 120


@dataclass(frozen=True, slots=True)
class ExcelResponse:
    """Én række fra den delte Forms-projektmappe."""

    response_id: int
    excel_row_number: int
    data: dict[str, Any]


class FormsExcelWorkbook:
    """Klient til den online/delte Excel-fil, som er knyttet til Forms.

    Objektet henter altid den aktuelle onlinefil. Ved sletning bruges filens
    eTag, så en samtidig ændring stopper uploaden i stedet for at overskrives.
    """

    def __init__(
        self,
        share_url: str,
        *,
        sheet_name: str = "Sheet1",
        id_column: str = "Id",
        credential_name: str = "API_SHAREPOINT",
        timeout_seconds: int = DEFAULT_TIMEOUT_SECONDS,
    ) -> None:
        self.share_url = _required_text(share_url, "share_url")
        self.sheet_name = _required_text(sheet_name, "sheet_name")
        self.id_column = _required_text(id_column, "id_column")
        self.credential_name = _required_text(
            credential_name,
            "credential_name",
        )
        if (
            isinstance(timeout_seconds, bool)
            or not isinstance(timeout_seconds, int)
            or timeout_seconds <= 0
        ):
            raise ValueError("timeout_seconds skal være et positivt heltal.")
        self.timeout_seconds = timeout_seconds
        self._token: str | None = None
        self._token_expires_at = 0.0

    def get_responses(self) -> list[ExcelResponse]:
        """Henter onlinefilen og returnerer alle svarrækker."""
        file_data, _ = self._download_current_workbook()
        responses = self._read_responses(file_data)

        logger.info("Excel-filen indeholder %s svarrækker.", len(responses))
        for response in responses:
            logger.info(
                "Excel-svar id=%s række=%s data=%s",
                response.response_id,
                response.excel_row_number,
                response.data,
            )
        return responses

    def get_response_by_id(self, response_id: int) -> ExcelResponse | None:
        """Returnerer den aktuelle Excel-række med det angivne Id."""
        validated_id = _validate_response_id(response_id)
        matches = [
            item
            for item in self.get_responses()
            if item.response_id == validated_id
        ]

        if len(matches) > 1:
            raise RuntimeError(
                f"Excel-filen indeholder flere rækker med Id {validated_id}."
            )
        return matches[0] if matches else None

    def delete_response_by_id(self, response_id: int) -> ExcelResponse:
        """Sletter præcis én Excel-række via kolonnen Id og uploader filen.

        Output:
            Den ExcelResponse, som blev slettet.
        """
        validated_id = _validate_response_id(response_id)
        file_data, metadata = self._download_current_workbook()
        updated_data, deleted_response = self._delete_row_in_workbook(
            workbook_bytes=file_data,
            response_id=validated_id,
        )
        self._upload_workbook(
            workbook_bytes=updated_data,
            metadata=metadata,
        )

        # Verificér mod en frisk download.
        if self.get_response_by_id(validated_id) is not None:
            raise RuntimeError(
                f"Excel-rækken med Id {validated_id} findes stadig efter upload."
            )

        logger.info("Excel-rækken med Id %s blev slettet.", validated_id)
        return deleted_response

    def _download_current_workbook(self) -> tuple[bytes, dict[str, Any]]:
        share_id = _encode_share_url(self.share_url)
        headers = self._headers()

        metadata_response = requests.get(
            f"{GRAPH_BASE_URL}/shares/{share_id}/driveItem",
            headers=headers,
            timeout=self.timeout_seconds,
        )
        _raise_graph_error(
            metadata_response,
            "Den delte Excel-fil kunne ikke findes",
        )
        metadata = metadata_response.json()

        content_response = requests.get(
            f"{GRAPH_BASE_URL}/shares/{share_id}/driveItem/content",
            headers=headers,
            timeout=self.timeout_seconds,
            allow_redirects=True,
        )
        _raise_graph_error(
            content_response,
            "Den delte Excel-fil kunne ikke downloades",
        )
        if not content_response.content:
            raise RuntimeError("Microsoft Graph returnerede en tom Excel-fil.")

        logger.info("Hentede onlinefilen: %s", metadata.get("name"))
        return content_response.content, metadata

    def _read_responses(self, workbook_bytes: bytes) -> list[ExcelResponse]:
        workbook = load_workbook(
            io.BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
        )
        try:
            sheet, headers, id_index = self._get_sheet_structure(workbook)
            responses: list[ExcelResponse] = []

            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                if not values or all(value is None for value in values):
                    continue

                raw_id = values[id_index] if id_index < len(values) else None
                response_id = _coerce_response_id(raw_id, row_number)
                row_data = {
                    header: values[index] if index < len(values) else None
                    for index, header in enumerate(headers)
                    if header
                }
                responses.append(
                    ExcelResponse(
                        response_id=response_id,
                        excel_row_number=row_number,
                        data=row_data,
                    )
                )
            return responses
        finally:
            workbook.close()

    def _delete_row_in_workbook(
        self,
        *,
        workbook_bytes: bytes,
        response_id: int,
    ) -> tuple[bytes, ExcelResponse]:
        workbook = load_workbook(io.BytesIO(workbook_bytes))
        try:
            sheet, headers, id_index = self._get_sheet_structure(workbook)
            matches: list[ExcelResponse] = []

            for row_number in range(2, sheet.max_row + 1):
                raw_id = sheet.cell(row=row_number, column=id_index + 1).value
                if raw_id is None:
                    continue
                current_id = _coerce_response_id(raw_id, row_number)
                if current_id != response_id:
                    continue

                row_data = {
                    header: sheet.cell(row=row_number, column=index + 1).value
                    for index, header in enumerate(headers)
                    if header
                }
                matches.append(
                    ExcelResponse(
                        response_id=current_id,
                        excel_row_number=row_number,
                        data=row_data,
                    )
                )

            if not matches:
                raise RuntimeError(
                    f"Excel-filen indeholder ingen række med Id {response_id}."
                )
            if len(matches) > 1:
                raise RuntimeError(
                    f"Excel-filen indeholder flere rækker med Id {response_id}."
                )

            deleted_response = matches[0]
            deleted_row = deleted_response.excel_row_number
            sheet.delete_rows(deleted_row, 1)
            self._adjust_tables_after_row_delete(sheet, deleted_row)

            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue(), deleted_response
        finally:
            workbook.close()

    def _get_sheet_structure(self, workbook):
        if self.sheet_name not in workbook.sheetnames:
            raise RuntimeError(
                f"Arket {self.sheet_name!r} findes ikke. "
                f"Tilgængelige ark: {workbook.sheetnames!r}."
            )
        sheet = workbook[self.sheet_name]
        headers = [
            str(cell.value).strip() if cell.value is not None else ""
            for cell in sheet[1]
        ]
        if self.id_column not in headers:
            raise RuntimeError(
                f"Kolonnen {self.id_column!r} findes ikke. Kolonner: {headers!r}."
            )
        return sheet, headers, headers.index(self.id_column)

    @staticmethod
    def _adjust_tables_after_row_delete(sheet, deleted_row: int) -> None:
        """Reducerer tabellens ref, når en datarække er slettet."""
        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if not min_row < deleted_row <= max_row:
                continue
            new_max_row = max(max_row - 1, min_row)
            table.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{new_max_row}"
            )

    def _upload_workbook(
        self,
        *,
        workbook_bytes: bytes,
        metadata: dict[str, Any],
    ) -> None:
        drive_id = ((metadata.get("parentReference") or {}).get("driveId"))
        item_id = metadata.get("id")
        etag = metadata.get("eTag")

        if not drive_id or not item_id:
            raise RuntimeError(
                "Graph-metadata mangler driveId eller item-id til upload."
            )

        headers = self._headers()
        headers["Content-Type"] = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        if etag:
            headers["If-Match"] = etag

        response = requests.put(
            f"{GRAPH_BASE_URL}/drives/{drive_id}/items/{item_id}/content",
            headers=headers,
            data=workbook_bytes,
            timeout=self.timeout_seconds,
        )
        _raise_graph_error(
            response,
            "Den opdaterede Excel-fil kunne ikke uploades",
        )

    def _headers(self) -> dict[str, str]:
        return {"Authorization": f"Bearer {self._get_token()}"}

    def _get_token(self) -> str:
        if self._token and time.time() < self._token_expires_at - 300:
            return self._token

        AutomationServer.from_environment()
        credential = Credential.get_credential(self.credential_name)
        config = credential.data or {}
        tenant_id = _required_text(config.get("tenant_id"), "tenant_id")
        client_id = _required_text(config.get("client_id"), "client_id")
        client_secret = _required_text(credential.password, "client_secret")

        response = requests.post(
            f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
            data={
                "grant_type": "client_credentials",
                "client_id": client_id,
                "client_secret": client_secret,
                "scope": GRAPH_SCOPE,
            },
            timeout=self.timeout_seconds,
        )
        _raise_graph_error(response, "Microsoft Graph-token kunne ikke hentes")
        token_data = response.json()
        self._token = _required_text(token_data.get("access_token"), "access_token")
        self._token_expires_at = time.time() + int(
            token_data.get("expires_in", 3600)
        )
        return self._token


def _validate_response_id(response_id: int) -> int:
    if isinstance(response_id, bool) or not isinstance(response_id, int):
        raise TypeError("response_id skal være et heltal.")
    if response_id <= 0:
        raise ValueError("response_id skal være større end 0.")
    return response_id


def _coerce_response_id(raw_id: Any, row_number: int) -> int:
    try:
        return _validate_response_id(int(raw_id))
    except (TypeError, ValueError) as error:
        raise RuntimeError(
            f"Ugyldigt Id i Excel-række {row_number}: {raw_id!r}."
        ) from error


def _encode_share_url(url: str) -> str:
    encoded = base64.urlsafe_b64encode(url.encode("utf-8")).decode("ascii")
    return "u!" + encoded.rstrip("=")


def _raise_graph_error(response: requests.Response, message: str) -> None:
    if response.ok:
        return
    try:
        error_data = response.json()
    except ValueError:
        error_data = {}
    graph_error = error_data.get("error") or {}
    raise RuntimeError(
        f"{message}. HTTP {response.status_code}. "
        f"Graph-fejl: {graph_error.get('code', 'ukendt_kode')}: "
        f"{graph_error.get('message', response.text or 'intet svar')}"
    )


def _required_text(value: Any, field_name: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise RuntimeError(f"{field_name} mangler eller er tom.")
    return value.strip()


__all__ = [
    "ExcelResponse",
    "FormsExcelWorkbook",
]
